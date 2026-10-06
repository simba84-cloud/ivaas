"""The Liquid brand on the reports people download: the PDFs, the Excel workbooks, and
the file names. The CSVs stay plain, header row first, because spreadsheets and ERP
imports read them; their names carry the brand instead.

The palette is the portal's (web/src/index.css): navy for the brand, magenta as the
one accent, a navy tint for alternate rows. The logo is the portal's own file, kept
here so the API image has it.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, TableStyle

LOGO = Path(__file__).parent / "assets" / "logo-liquid.png"
LOGO_ASPECT = 768 / 224  # width over height of the file
ORG = "Liquid Intelligent Technologies"
PRODUCT = "IVaaS crate counting"

NAVY, NAVY_DEEP, MAGENTA = "273C87", "1D2D66", "C8187D"
TINT, RULE, MUTED = "E9ECF5", "D5D9E6", "5B6275"


# file names -------------------------------------------------------------------------


def filename(site: str, *parts: str, ext: str) -> str:
    """liquid-ivaas-bakery-industrial-site-2026-10-01.pdf"""
    words = [site or "site", *parts]
    slug = re.sub(r"[^a-z0-9]+", "-", "-".join(words).lower()).strip("-")
    return f"liquid-ivaas-{slug}.{ext}"


# PDF --------------------------------------------------------------------------------


def pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    navy = colors.HexColor(f"#{NAVY}")
    return {
        "h1": ParagraphStyle("h1", parent=base["Heading1"], textColor=navy, spaceAfter=2),
        "h2": ParagraphStyle(
            "h2", parent=base["Heading3"], textColor=navy, spaceBefore=8, spaceAfter=3
        ),
        "body": base["BodyText"],
    }


#: header row navy with white type, alternate rows tinted, hairline rules
GRID = TableStyle(
    [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor(f"#{TINT}")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor(f"#{RULE}")),
        ("BOX", (0, 0), (-1, -1), 0.25, colors.HexColor(f"#{RULE}")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
)


def pdf_document(buf: io.BytesIO, pagesize: tuple[float, float], title: str) -> Any:
    """A document whose every page carries the logo and the report's name at the top,
    and the brand and page number at the foot (see `pdf_page`)."""
    return SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=26 * mm,
        bottomMargin=16 * mm,
        title=title,
        author=f"{PRODUCT}, {ORG}",
        creator=f"{PRODUCT}, {ORG}",
    )


def pdf_page(heading: str):
    """The onPage callback: logo top left, `heading` top right, a magenta rule under
    them; the brand bottom left and the page number bottom right."""

    def draw(canvas: Any, doc: Any) -> None:
        width, height = doc.pagesize
        left, right = doc.leftMargin, width - doc.rightMargin
        logo_h = 9 * mm
        top = height - 9 * mm
        canvas.saveState()
        canvas.drawImage(
            str(LOGO),
            left,
            top - logo_h,
            width=logo_h * LOGO_ASPECT,
            height=logo_h,
            mask="auto",
        )
        canvas.setFillColor(colors.HexColor(f"#{NAVY}"))
        canvas.setFont("Helvetica-Bold", 10)
        canvas.drawRightString(right, top - 4 * mm, heading)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawRightString(right, top - 8 * mm, PRODUCT)
        canvas.setStrokeColor(colors.HexColor(f"#{MAGENTA}"))
        canvas.setLineWidth(1.2)
        canvas.line(left, top - logo_h - 3 * mm, right, top - logo_h - 3 * mm)

        canvas.setStrokeColor(colors.HexColor(f"#{RULE}"))
        canvas.setLineWidth(0.5)
        canvas.line(left, 11 * mm, right, 11 * mm)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.setFont("Helvetica", 7.5)
        canvas.drawString(left, 7 * mm, f"{PRODUCT}  ·  {ORG}")
        canvas.drawRightString(right, 7 * mm, f"Page {doc.page}")
        canvas.restoreState()

    return draw


# Excel ------------------------------------------------------------------------------

_HEAD_FONT = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor=NAVY)
_ZEBRA = PatternFill("solid", fgColor=TINT)
_RULE = Border(bottom=Side(style="thin", color=RULE))
#: rows the logo sits in at the top of every sheet
_LOGO_ROWS = 3


class Sheet:
    """One branded worksheet: the logo, the title and a line under it, then sections
    added top to bottom. Numbers stay numbers, so the workbook can be summed and
    filtered; a figure the records cannot support is left empty, never 0."""

    def __init__(self, wb: Workbook, name: str, title: str, subtitle: str) -> None:
        self.ws = wb.create_sheet(name[:31])
        self.ws.sheet_view.showGridLines = False
        for r in range(1, _LOGO_ROWS + 1):
            self.ws.row_dimensions[r].height = 15
        logo = XlImage(str(LOGO))
        logo.height = 40
        logo.width = round(40 * LOGO_ASPECT)
        self.ws.add_image(logo, "A1")
        self.row = _LOGO_ROWS + 1
        self.ws.cell(self.row, 1, title).font = Font(bold=True, size=14, color=NAVY)
        self.row += 1
        self.ws.cell(self.row, 1, subtitle).font = Font(size=9, color=MUTED)
        self.row += 2
        self._widths: dict[int, int] = {}

    def heading(self, text: str) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font = Font(bold=True, size=11, color=NAVY)
        cell.border = Border(bottom=Side(style="medium", color=MAGENTA))
        self.row += 1

    def note(self, text: str) -> None:
        self.ws.cell(self.row, 1, text).font = Font(italic=True, size=9, color=MUTED)
        self.row += 1

    def gap(self) -> None:
        self.row += 1

    def table(
        self,
        headers: list[str],
        rows: list[list[Any]],
        *,
        formats: dict[int, str] | None = None,
        filterable: bool = False,
    ) -> None:
        """`formats` maps a column index to an Excel number format, e.g. "0.0%"."""
        formats = formats or {}
        top = self.row
        for i, h in enumerate(headers, start=1):
            cell = self.ws.cell(top, i, h)
            cell.font, cell.fill = _HEAD_FONT, _HEAD_FILL
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            self._fit(i, h)
        for n, values in enumerate(rows, start=1):
            for i, v in enumerate(values, start=1):
                cell = self.ws.cell(top + n, i, None if v == "" else v)
                cell.border = _RULE
                if n % 2 == 0:
                    cell.fill = _ZEBRA
                if i - 1 in formats and v not in ("", None):
                    cell.number_format = formats[i - 1]
                elif isinstance(v, int) and not isinstance(v, bool):
                    cell.number_format = "#,##0"
                self._fit(i, v)
        if filterable:
            last = get_column_letter(len(headers))
            self.ws.auto_filter.ref = f"A{top}:{last}{top + max(len(rows), 1)}"
            self.ws.freeze_panes = self.ws.cell(top + 1, 1)
        self.row = top + len(rows) + 2

    def _fit(self, col: int, value: Any) -> None:
        text = "" if value is None else str(value)
        width = min(max(len(text) + 2, 9), 48)
        if width > self._widths.get(col, 0):
            self._widths[col] = width
            self.ws.column_dimensions[get_column_letter(col)].width = width


def workbook(title: str) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = title
    wb.properties.creator = f"{PRODUCT}, {ORG}"
    return wb


def xlsx_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
