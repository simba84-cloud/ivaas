"""The crate tally sheet workbook, made for a tenant: the paper form to print, the two
entry sheets the office types into, and the lists behind their drop-downs.

It is the template the POC started with (Bakers_Inn_Crate_Tally_Sheet_Template.xlsx),
remade so that it cannot drift from the importer: the entry sheets carry exactly the
columns `tally_csv` reads, the EXAMPLE rows are the ones it skips, and the bays in
the drop-down are the tenant's own. The paper form and the instructions carry the
Liquid brand (`branding`); the entry sheets keep their header row in row 1, because
that is the row the importer reads.
"""

from __future__ import annotations

import re
from datetime import date, datetime, time

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ivaas.adapters import branding
from ivaas.adapters.branding import LOGO, LOGO_ASPECT, MAGENTA, MUTED, NAVY, ORG, PRODUCT, TINT

VERSION = "0.2"
SHEET_COLUMNS = [
    "sheet_id", "date", "bay", "truck_plate", "direction", "start_time", "end_time",
    "route_driver", "pages", "stack_count", "total_crates", "total_on_paper", "check",
    "counted_by", "verified_by", "entered_by", "notes",
]  # fmt: skip
STACK_COLUMNS = ["sheet_id", "line_no", "crates", "note"]
NOTE_CODES = [
    ("P", "Partial or broken stack"),
    ("X", "Crossed back or removed (write a negative count)"),
    ("D", "Damaged crates"),
    ("?", "Unsure"),
]
#: room for a busy month of sheets, and their lines
SHEET_ROWS, STACK_ROWS = 500, 5000
#: lines on one page of the paper form: three columns of twenty
LINES_PER_COLUMN = 20

_THIN = Side(style="thin", color="9AA1B8")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor=NAVY)
_INPUT = PatternFill("solid", fgColor=TINT)
_EXAMPLE = PatternFill("solid", fgColor="EEEEEE")
_EXAMPLE_FONT = Font(italic=True, color=MUTED)
_LABEL = Font(bold=True, color=NAVY, size=10)


def prefix_of(tenant: str) -> str:
    """The Sheet ID prefix: the tenant's initials, BI for Bakers Inn."""
    words = re.findall(r"[A-Za-z0-9]+", tenant)
    return "".join(w[0] for w in words).upper()[:4] or "TS"


def _logo(ws, anchor: str, height: int = 40) -> None:
    img = XlImage(str(LOGO))
    img.height, img.width = height, round(height * LOGO_ASPECT)
    ws.add_image(img, anchor)


def _footer(ws) -> None:
    ws.oddFooter.left.text = f"{PRODUCT} · {ORG}"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8


def _a4(ws, *, landscape: bool = False) -> None:
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0 if landscape else 1
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top, ws.page_margins.bottom = 0.4, 0.6


def _instructions(wb: Workbook, tenant: str, prefix: str, target: float, made: date) -> None:
    ws = wb.create_sheet("Instructions")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 15
    _logo(ws, "A1")
    ws["A4"] = f"{tenant} crate tally sheet: instructions"
    ws["A4"].font = Font(bold=True, size=14, color=NAVY)
    ws["A5"] = (
        f"Ground truth for the {PRODUCT} proof of concept. Template v{VERSION}, made "
        f"{made:%d %B %Y} by the portal for {tenant}."
    )
    ws["A5"].font = Font(size=9, color=MUTED)

    sample = f"{prefix}-{made:%Y%m%d}-B1-001"
    sections: list[tuple[str, list[tuple[str, str]]]] = [
        (
            "This workbook",
            [
                (
                    "Tally Sheet",
                    "The paper form (A4). Print one per truck per session: one "
                    "LOAD or one RETURN. Use another page if a truck has more than 60 dollies "
                    "or stacks.",
                ),
                (
                    "Entry - Sheets",
                    "Office entry: one row per paper tally sheet. Enter each "
                    "day's sheets before the end of the next working morning.",
                ),
                (
                    "Entry - Stacks",
                    "Office entry: one row per dolly or stack line on the paper "
                    "sheet, with the same Sheet ID.",
                ),
                ("Lists", "The drop-down values: your bays, the directions and the note codes."),
            ],
        ),
        (
            "Legend",
            [
                (
                    "Shaded cells",
                    "Write or type here. Other cells are labels or formulas: "
                    "leave them as they are.",
                ),
                (
                    "Grey, italic rows",
                    "EXAMPLE rows showing the format. The portal skips any "
                    "Sheet ID starting with EXAMPLE, so they can stay.",
                ),
            ],
        ),
        (
            "How to count",
            [
                (
                    "1. Before loading",
                    "Fill in the header: date, bay, truck plate, direction, "
                    f"start time and your name. Sheet ID = {prefix}-YYYYMMDD-<bay>-<3-digit "
                    f"number>, for example {sample}.",
                ),
                (
                    "2. During loading",
                    "Count the crates on each dolly or stack as it crosses "
                    "into the truck (or out of it for a RETURN). One line per dolly or stack, in "
                    "order.",
                ),
                (
                    "3. Anything unusual",
                    "Use the Note column: "
                    + "; ".join(f"{c} = {d.lower()}" for c, d in NOTE_CODES)
                    + ".",
                ),
                (
                    "4. After loading",
                    "Write the end time and the truck's total crates, then "
                    "sign. The supervisor checks the total and countersigns.",
                ),
                (
                    "5. Independence",
                    "Do not look at the IVaaS screen or reports while counting "
                    "or entering. The portal keeps the AI count back until the sheet is saved.",
                ),
            ],
        ),
        (
            "Entering and uploading",
            [
                (
                    "Entry - Sheets",
                    "One row per paper sheet. stack_count and total_crates are "
                    "worked out from Entry - Stacks; type the handwritten total into "
                    "total_on_paper. check shows MISMATCH when they differ: fix the typing "
                    "before uploading.",
                ),
                (
                    "Entry - Stacks",
                    "One row per line on the paper sheet: sheet_id, line_no, crates, note.",
                ),
                (
                    "Upload",
                    "In the portal, open Tally sheets, choose the bay, and upload this "
                    "workbook as it is (the Workbook box). Both entry sheets are read from it; "
                    "nothing needs saving as CSV. Keep the paper originals for 90 days.",
                ),
                (
                    "Matching",
                    "Each sheet is matched to a counted load by truck plate, bay and "
                    "time. A sheet that matches no load is listed in the portal for review.",
                ),
            ],
        ),
        (
            "Accuracy",
            [
                (
                    "Per load",
                    "accuracy = 1 - |AI count - tally count| / tally count, on the AI "
                    f"count and never on a correction. The POC target is {target:.0%} or better.",
                ),
            ],
        ),
    ]
    row = 7
    for heading, items in sections:
        cell = ws.cell(row, 1, heading.upper())
        cell.font = Font(bold=True, color=NAVY)
        cell.border = Border(bottom=Side(style="medium", color=MAGENTA))
        row += 1
        for label, text in items:
            ws.cell(row, 1, label).font = Font(bold=True, size=10)
            body = ws.cell(row, 2, text)
            body.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row, 1).alignment = Alignment(vertical="top")
            ws.row_dimensions[row].height = 15 * max(1, -(-len(text) // 105))
            row += 1
        row += 1
    _a4(ws)
    ws.page_setup.fitToHeight = 0
    _footer(ws)


def _form(wb: Workbook, tenant: str, prefix: str) -> None:
    ws = wb.create_sheet("Tally Sheet")
    ws.sheet_view.showGridLines = False
    for col, width in zip("ABCDEFGHIJK", (6, 9, 10, 2, 6, 9, 10, 2, 6, 9, 10), strict=True):
        ws.column_dimensions[col].width = width
    for r in (1, 2, 3):
        ws.row_dimensions[r].height = 15
    _logo(ws, "A1", height=36)
    ws.merge_cells("F1:K2")
    ws["F1"] = "CRATE TALLY SHEET"
    ws["F1"].font = Font(bold=True, size=15, color=NAVY)
    ws["F1"].alignment = Alignment(horizontal="right", vertical="center")
    ws.merge_cells("F3:K3")
    ws["F3"] = f"{tenant} · {PRODUCT}"
    ws["F3"].font = Font(size=8, color=MUTED)
    ws["F3"].alignment = Alignment(horizontal="right")
    for col in range(1, 12):
        ws.cell(4, col).border = Border(top=Side(style="medium", color=MAGENTA))
    ws.merge_cells("A5:K5")
    ws["A5"] = (
        "One sheet per truck per session. Count each dolly or stack as it crosses. "
        "Do NOT look at the IVaaS screen while counting."
    )
    ws["A5"].font = Font(italic=True, size=9, color=MUTED)

    def field(row: int, label: str, cols: str, value_cols: str, hint: str = "") -> None:
        a, b = cols.split(":")
        ws.merge_cells(f"{a}{row}:{b}{row}")
        ws[f"{a}{row}"] = label
        ws[f"{a}{row}"].font = _LABEL
        c, d = value_cols.split(":")
        ws.merge_cells(f"{c}{row}:{d}{row}")
        first, last = ord(c), ord(d)
        for col in range(first, last + 1):
            cell = ws[f"{chr(col)}{row}"]
            cell.fill, cell.border = _INPUT, _BOX
        if hint:
            ws[f"{c}{row}"] = hint
            ws[f"{c}{row}"].font = Font(size=9, color=MUTED)
        ws.row_dimensions[row].height = 22

    field(7, "Sheet ID", "A:B", "C:F", f"{prefix}-YYYYMMDD-___-___")
    field(7, "Date", "G:H", "I:K")
    field(8, "Bay", "A:B", "C:F")
    field(8, "Direction", "G:H", "I:K", "☐ LOAD     ☐ RETURN")
    field(9, "Truck plate", "A:B", "C:F")
    field(9, "Route / driver", "G:H", "I:K")
    field(10, "Start time", "A:B", "C:F")
    field(10, "End time", "G:H", "I:K")
    field(11, "Counted by", "A:B", "C:F")
    field(11, "Signature", "G:H", "I:K")
    field(12, "Page", "A:B", "C:F", "____ of ____")

    top = 14
    blocks = [("A", "B", "C"), ("E", "F", "G"), ("I", "J", "K")]
    for n, (no, crates, note) in enumerate(blocks):
        for col, title in ((no, "No."), (crates, "Crates"), (note, "Note")):
            cell = ws[f"{col}{top}"]
            cell.value, cell.font, cell.fill = title, _HEAD, _HEAD_FILL
            cell.alignment = Alignment(horizontal="center")
        for i in range(LINES_PER_COLUMN):
            r = top + 1 + i
            ws[f"{no}{r}"] = n * LINES_PER_COLUMN + i + 1
            ws[f"{no}{r}"].font = Font(size=9, color=MUTED)
            ws[f"{no}{r}"].alignment = Alignment(horizontal="center")
            for col in (no, crates, note):
                ws[f"{col}{r}"].border = _BOX
            for col in (crates, note):
                ws[f"{col}{r}"].fill = _INPUT
            ws.row_dimensions[r].height = 19
    first, last = top + 1, top + LINES_PER_COLUMN
    counted = ",".join(f"{c}{first}:{c}{last}" for _, c, _ in blocks)

    end = last + 2
    totals = [
        (
            "Dollies / stacks on this page",
            f'=IF(COUNT({counted})=0,"",COUNT({counted}))',
            "Worked out when typed in; handwrite it on paper.",
        ),
        (
            "Total crates on this page",
            f'=IF(COUNT({counted})=0,"",SUM({counted}))',
            "Worked out when typed in; handwrite it on paper.",
        ),
        ("TOTAL CRATES FOR THE TRUCK (all pages)", None, "On the last page only."),
    ]
    for i, (label, formula, hint) in enumerate(totals):
        r = end + i
        ws.merge_cells(f"A{r}:E{r}")
        ws[f"A{r}"] = label
        ws[f"A{r}"].font = _LABEL
        ws.merge_cells(f"F{r}:G{r}")
        for col in "FG":
            ws[f"{col}{r}"].border, ws[f"{col}{r}"].fill = _BOX, _INPUT
        if formula:
            ws[f"F{r}"] = formula
        ws[f"F{r}"].alignment = Alignment(horizontal="center")
        ws.merge_cells(f"H{r}:K{r}")
        ws[f"H{r}"] = hint
        ws[f"H{r}"].font = Font(size=8, italic=True, color=MUTED)
        ws.row_dimensions[r].height = 22
    sign = end + 4
    field(sign, "Supervisor", "A:B", "C:F")
    field(sign, "Signature", "G:H", "I:K")
    field(sign + 1, "Entered by (office)", "A:B", "C:F")
    field(sign + 1, "Entry date", "G:H", "I:K")
    codes = sign + 3
    ws.merge_cells(f"A{codes}:K{codes}")
    ws[f"A{codes}"] = (
        "Note codes: "
        + " · ".join(f"{c} = {d.lower()}" for c, d in NOTE_CODES)
        + ". Keep this sheet for 90 days."
    )
    ws[f"A{codes}"].font = Font(size=8, color=MUTED)
    ws[f"A{codes}"].alignment = Alignment(wrap_text=True)
    ws.row_dimensions[codes].height = 24
    ws.print_area = f"A1:K{codes}"
    _a4(ws)
    _footer(ws)


def _entry_sheets(wb: Workbook, bays: list[str], made: date, prefix: str) -> None:
    ws = wb.create_sheet("Entry - Sheets")
    widths = (26, 12, 16, 13, 10, 10, 10, 22, 7, 11, 12, 14, 11, 14, 14, 14, 36)
    for i, (name, width) in enumerate(zip(SHEET_COLUMNS, widths, strict=True), start=1):
        cell = ws.cell(1, i, name)
        cell.font, cell.fill = _HEAD, _HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = width
    example = [
        f"EXAMPLE-{made:%Y%m%d}-B1-001", datetime.combine(made, time()), bays[0] if bays else "B1",
        "AEX 4821", "LOAD", time(6, 40), time(7, 5), "Route 7 / T. Moyo", 1, None, None, 124,
        None, "R. Ncube", "S. Dube", "P. Chuma", "Line 5: one dolly pulled back off the truck",
    ]  # fmt: skip
    for i, v in enumerate(example, start=1):
        if v is not None:
            ws.cell(2, i, v)
    stacks = "'Entry - Stacks'"
    span = STACK_ROWS + 1
    for r in range(2, SHEET_ROWS + 2):
        ws[f"J{r}"] = f'=IF($A{r}="","",COUNTIFS({stacks}!$A$2:$A${span},$A{r}))'
        ws[f"K{r}"] = (
            f'=IF($A{r}="","",SUMIFS({stacks}!$C$2:$C${span},{stacks}!$A$2:$A${span},$A{r}))'
        )
        ws[f"M{r}"] = f'=IF(OR($A{r}="",L{r}=""),"",IF(K{r}=L{r},"OK","MISMATCH"))'
        ws[f"B{r}"].number_format = "yyyy-mm-dd"
        ws[f"F{r}"].number_format = ws[f"G{r}"].number_format = "hh:mm"
    for col in range(1, len(SHEET_COLUMNS) + 1):
        c = ws.cell(2, col)
        c.fill, c.font = _EXAMPLE, _EXAMPLE_FONT
    last = SHEET_ROWS + 1
    ws.conditional_formatting.add(
        f"M2:M{last}",
        FormulaRule(
            formula=['$M2="MISMATCH"'],
            fill=PatternFill("solid", fgColor="FBE3E3"),
            font=Font(bold=True, color="D13B3B"),
        ),
    )
    rules = [
        DataValidation(type="list", formula1="=Lists!$B$2:$B$3", allow_blank=True),
        DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)"),
        DataValidation(type="time", operator="between", formula1="0", formula2="0.999988"),
        DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0"),
    ]
    rules[0].add(f"E2:E{last}")
    rules[1].add(f"B2:B{last}")
    rules[2].add(f"F2:G{last}")
    rules[3].add(f"I2:I{last}")
    rules[3].add(f"L2:L{last}")
    for dv in rules:
        dv.error = "That value is not allowed here: see the Instructions sheet."
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
    if bays:
        dv = DataValidation(type="list", formula1=f"=Lists!$A$2:$A${len(bays) + 1}")
        dv.add(f"C2:C{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(SHEET_COLUMNS))}{last}"
    _a4(ws, landscape=True)
    _footer(ws)


def _entry_stacks(wb: Workbook, made: date) -> None:
    ws = wb.create_sheet("Entry - Stacks")
    for i, (name, width) in enumerate(zip(STACK_COLUMNS, (26, 9, 9, 8), strict=True), start=1):
        cell = ws.cell(1, i, name)
        cell.font, cell.fill = _HEAD, _HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = width
    sid = f"EXAMPLE-{made:%Y%m%d}-B1-001"
    lines = [
        (1, 32, None),
        (2, 32, None),
        (3, 28, "P"),
        (4, 32, None),
        (5, -32, "X"),
        (6, 32, None),
    ]
    for r, (no, crates, note) in enumerate(lines, start=2):
        for col, v in enumerate((sid, no, crates, note), start=1):
            c = ws.cell(r, col, v)
            c.fill, c.font = _EXAMPLE, _EXAMPLE_FONT
    last = STACK_ROWS + 1
    line = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1")
    line.add(f"B2:B{last}")
    crates = DataValidation(type="whole", operator="between", formula1="-500", formula2="500")
    crates.add(f"C2:C{last}")
    note = DataValidation(type="list", formula1=f"=Lists!$C$2:$C${len(NOTE_CODES) + 1}")
    note.add(f"D2:D{last}")
    for dv in (line, crates, note):
        dv.allow_blank = True
        dv.error = "That value is not allowed here: see the Instructions sheet."
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
    ws.freeze_panes = "A2"
    _footer(ws)


def _lists(wb: Workbook, bays: list[str]) -> None:
    ws = wb.create_sheet("Lists")
    for i, (name, width) in enumerate(
        zip(("bay", "direction", "note_code", "meaning"), (24, 11, 11, 44), strict=True), start=1
    ):
        cell = ws.cell(1, i, name)
        cell.font, cell.fill = _HEAD, _HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = width
    for r, bay in enumerate(bays, start=2):
        ws.cell(r, 1, bay)
    for r, direction in enumerate(("LOAD", "RETURN"), start=2):
        ws.cell(r, 2, direction)
    for r, (code, meaning) in enumerate(NOTE_CODES, start=2):
        ws.cell(r, 3, code)
        ws.cell(r, 4, meaning)


def tally_template(tenant: str, bays: list[str], *, target: float, made: date) -> bytes:
    """The workbook for `tenant`, its `bays` in the drop-down. A tenant with no bays yet
    gets an empty bay list, not an invented one."""
    prefix = prefix_of(tenant)
    wb = branding.workbook(f"{tenant} crate tally sheet")
    _instructions(wb, tenant, prefix, target, made)
    _form(wb, tenant, prefix)
    _entry_sheets(wb, bays, made, prefix)
    _entry_stacks(wb, made)
    _lists(wb, bays)
    wb.active = 0
    return branding.xlsx_bytes(wb)
