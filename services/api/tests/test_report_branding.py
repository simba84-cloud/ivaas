"""Reports people download carry the Liquid brand: the logo and colours on the PDF and
the Excel workbook, the brand in every file name. The CSV stays plain for imports."""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime

from conftest import login, make_client
from openpyxl import load_workbook
from pypdf import PdfReader
from test_edge import bakers  # noqa: F401
from test_poc_report import DAYS, the_poc
from test_reports import DAY, day  # noqa: F401

from ivaas.adapters.branding import NAVY, filename

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _get(c, admin, bay, fmt, d=DAY):
    return c.get(
        "/api/v1/reports/daily",
        params={"site_id": bay["site_id"], "day": d, "format": fmt},
        headers=admin,
    )


def _table(ws, first: str):
    """The rows of the table whose header row starts with `first`, as dicts."""
    rows = list(ws.iter_rows(values_only=True))
    top = next(i for i, r in enumerate(rows) if r[0] == first)
    head = rows[top]
    out = []
    for r in rows[top + 1 :]:
        if r[0] is None:
            break
        out.append(dict(zip(head, r, strict=False)))
    return top + 1, out


def test_file_names_carry_the_brand():
    assert filename("Bakery Industrial Site", "2026-10-01", ext="pdf") == (
        "liquid-ivaas-bakery-industrial-site-2026-10-01.pdf"
    )
    assert filename("", "poc", "2026-10-01", "to", "2026-10-03", ext="xlsx") == (
        "liquid-ivaas-site-poc-2026-10-01-to-2026-10-03.xlsx"
    )


def test_every_pdf_page_has_the_logo_and_the_brand_at_the_foot(day):  # noqa: F811
    c, admin, bay, _ = day
    r = _get(c, admin, bay, "pdf")
    assert r.headers["content-disposition"] == (
        'attachment; filename="liquid-ivaas-bakery-industrial-site-2026-10-01.pdf"'
    )
    pdf = PdfReader(io.BytesIO(r.content))
    for n, page in enumerate(pdf.pages, start=1):
        assert len(page.images) >= 1  # the logo
        text = page.extract_text()
        assert "Liquid Intelligent Technologies" in text and f"Page {n}" in text
        assert "Daily report" in text
    assert "Liquid Intelligent Technologies" in pdf.metadata.author


def test_the_csv_stays_plain_under_a_branded_name(day):  # noqa: F811
    c, admin, bay, _ = day
    r = _get(c, admin, bay, "csv")
    assert r.headers["content-disposition"].endswith(
        'filename="liquid-ivaas-bakery-industrial-site-2026-10-01.csv"'
    )
    assert r.text.startswith("date,site,opened")  # the header row first, nothing above it
    assert len(list(csv.DictReader(io.StringIO(r.text)))) == 5


def test_the_workbook_has_the_logo_and_the_figures_as_numbers(day):  # noqa: F811
    c, admin, bay, _ = day
    r = _get(c, admin, bay, "xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"].endswith(
        'filename="liquid-ivaas-bakery-industrial-site-2026-10-01.xlsx"'
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Every load"]
    assert all(len(ws._images) == 1 for ws in wb.worksheets)  # the logo on each sheet

    summary = wb["Summary"]
    _, [totals] = _table(summary, "Loads")
    # the same figures as the PDF: dispatched 100 + 78 (corrected) + 12, returned 30
    assert totals["Crates dispatched"] == 190 and totals["Crates returned"] == 30
    _, [acc] = _table(summary, "Verified loads")
    assert round(acc["Mean accuracy"], 4) == 0.9796

    loads = wb["Every load"]
    top, rows = _table(loads, "Opened")
    head = loads.cell(top, 1)
    assert head.font.b and head.fill.fgColor.rgb.endswith(NAVY)
    assert loads.freeze_panes == f"A{top + 1}" and loads.auto_filter.ref
    first, corrected = rows[0], rows[1]
    assert first["AI count"] == 100 and first["Tally"] == 98
    assert round(first["Accuracy"], 4) == 0.9796
    assert loads.cell(top + 1, 10).number_format == "0.0%"
    # nothing measured is left empty, never written as 0
    assert first["Corrected"] is None and corrected["Tally"] is None
    assert corrected["Corrected"] == 78 and corrected["Of record"] == 78


def test_a_workbook_for_a_day_without_tally_sheets_says_so_rather_than_zero():
    with make_client() as c:
        admin = login(c, "admin")
        site = c.get("/api/v1/sites", headers=admin).json()[0]["id"]
        r = c.get(
            "/api/v1/reports/daily",
            params={"site_id": site, "day": "2026-01-01", "format": "xlsx"},
            headers=admin,
        )
        wb = load_workbook(io.BytesIO(r.content))
        values = [v for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for v in row]
        assert any(isinstance(v, str) and "no accuracy figure" in v for v in values)
        assert "Mean accuracy" not in values
        assert any(isinstance(v, str) and "No loads were counted" in v for v in values)


def test_the_filed_report_keeps_a_workbook_too(day):  # noqa: F811
    c, admin, _, clock = day
    container = c.app.state.container
    from ivaas.domain.tenancy import BAKERS_INN_ID
    from ivaas.tenancy import tenant_context

    async def file():
        with tenant_context(BAKERS_INN_ID):
            return await (await container.file_daily_reports_uc())()

    clock.at = datetime(2026, 10, 2, 4, 30, tzinfo=UTC)
    c.portal.call(file)
    [listed] = c.get("/api/v1/reports", headers=admin).json()
    assert listed["file_stem"] == "liquid-ivaas-bakery-industrial-site-2026-10-01"
    wb = load_workbook(io.BytesIO(c.get(listed["xlsx_url"]).content))
    assert wb.sheetnames == ["Summary", "Every load"]


def test_the_poc_workbook_carries_the_same_figures_as_the_pdf(bakers):  # noqa: F811
    c, admin, bay, clock = bakers
    the_poc(c, admin, bay, clock)
    params = {"site_id": bay["site_id"], **DAYS, "baseline_minutes": 25, "format": "xlsx"}
    r = c.get("/api/v1/reports/poc", params=params, headers=admin)
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert 'filename="liquid-ivaas-' in r.headers["content-disposition"]
    assert "-poc-" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Result", "Every load"]
    result = [v for row in wb["Result"].iter_rows(values_only=True) for v in row]
    assert "Result: FAIL" in result and "66.7%" in result
    assert any(isinstance(v, str) and "not priced" in v for v in result)
    _, rows = _table(wb["Every load"], "Day")
    assert len(rows) == 4
    wrong = next(x for x in rows if x["Read by camera"] == "ABC 1008")
    assert wrong["Tally"] == 60
    pdf = c.get("/api/v1/reports/poc", params={**params, "format": "pdf"}, headers=admin)
    assert all(len(p.images) >= 1 for p in PdfReader(io.BytesIO(pdf.content)).pages)
