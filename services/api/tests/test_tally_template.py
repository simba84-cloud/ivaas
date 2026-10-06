"""The tally sheet workbook the portal hands out: Liquid-branded, this tenant's bays in
the drop-down, and entry sheets that are exactly what the upload reads, so a filled-in
copy uploads as it is."""

from __future__ import annotations

import io
from datetime import datetime, time

from openpyxl import load_workbook
from test_sheet_uploads import XLSX
from test_tally_api import bay, local_window, roles, session, truck_session  # noqa: F401

from ivaas.adapters.tally_template import SHEET_COLUMNS, STACK_COLUMNS, prefix_of


def _template(anon, headers):
    r = anon.get("/api/v1/tally/template", headers=headers)
    assert r.status_code == 200, r.text
    return r


def _upload(anon, bay, headers, raw):  # noqa: F811
    return anon.post(
        f"/api/v1/tally/import?bay_id={bay['id']}",
        files={"sheets": ("tally.xlsx", raw, XLSX)},
        headers=headers,
    )


def test_the_template_is_branded_and_made_for_the_tenant(anon, roles):  # noqa: F811
    r = _template(anon, roles["operator"])
    assert r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"].endswith(
        'filename="liquid-ivaas-bakers-inn-tally-sheet.xlsx"'
    )
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == [
        "Instructions",
        "Tally Sheet",
        "Entry - Sheets",
        "Entry - Stacks",
        "Lists",
    ]
    assert len(wb["Instructions"]._images) == 1 and len(wb["Tally Sheet"]._images) == 1
    assert "Liquid Intelligent Technologies" in wb["Tally Sheet"].oddFooter.left.text
    text = " ".join(
        str(v) for row in wb["Instructions"].iter_rows(values_only=True) for v in row if v
    )
    # the instructions describe the portal as it is: upload the workbook, no CSV step
    assert "upload this workbook as it is" in text and "Save As" not in text
    assert "BI-" in text  # Bakers Inn's sheet IDs
    # the paper form prints on one A4 page
    form = wb["Tally Sheet"]
    assert str(form.page_setup.paperSize) == str(form.PAPERSIZE_A4) and form.print_area


def test_the_entry_sheets_are_what_the_upload_reads(anon, roles):  # noqa: F811
    wb = load_workbook(io.BytesIO(_template(anon, roles["operator"]).content))
    sheets, stacks = wb["Entry - Sheets"], wb["Entry - Stacks"]
    assert [c.value for c in sheets[1]] == SHEET_COLUMNS
    assert [c.value for c in stacks[1]][: len(STACK_COLUMNS)] == STACK_COLUMNS
    assert sheets["A2"].value.startswith("EXAMPLE-") and stacks["A2"].value.startswith("EXAMPLE-")
    assert sheets["M3"].value.startswith("=IF(")  # the MISMATCH check is a formula
    # the drop-downs: the tenant's bays (never invented), the directions, the note codes
    bays = [b["name"] for b in anon.get("/api/v1/bays", headers=roles["viewer"]).json()]
    lists = wb["Lists"]
    assert [lists.cell(r, 1).value for r in range(2, len(bays) + 2)] == bays
    assert lists.cell(len(bays) + 2, 1).value is None
    assert [lists["B2"].value, lists["B3"].value] == ["LOAD", "RETURN"]
    guarded = {str(dv.sqref) for dv in sheets.data_validations.dataValidation}
    assert any(s.startswith("C2:") for s in guarded) and any(s.startswith("E2:") for s in guarded)


def test_the_untouched_template_says_only_the_examples_are_filled_in(anon, roles, bay):  # noqa: F811
    raw = _template(anon, roles["operator"]).content
    r = _upload(anon, bay, roles["operator"], raw)
    assert r.status_code == 422 and "Only the EXAMPLE rows are filled in" in r.text


def test_a_filled_in_template_uploads_as_it_is_and_reconciles(anon, roles, bay):  # noqa: F811
    op = roles["operator"]
    sid = truck_session(anon, bay, op, crates=(15, 15, 14))  # the AI saw 44
    wb = load_workbook(io.BytesIO(_template(anon, op).content))
    day, start, end = local_window()
    sheets, stacks = wb["Entry - Sheets"], wb["Entry - Stacks"]
    hm = lambda s: time(*map(int, s.split(":")))  # noqa: E731
    row = ["BI-1", datetime.fromisoformat(day), bay["name"], "ABE 2437", "LOAD", hm(start), hm(end)]
    for col, v in enumerate(row, start=1):
        sheets.cell(3, col, v)
    sheets["L3"] = 45  # total_on_paper
    for r, crates in enumerate((15, 15, 15), start=8):  # below the six example lines
        stacks.cell(r, 1, "BI-1")
        stacks.cell(r, 2, r - 7)
        stacks.cell(r, 3, crates)
    buf = io.BytesIO()
    wb.save(buf)
    r = _upload(anon, bay, op, buf.getvalue())
    assert r.status_code == 200, r.text
    (saved,) = r.json()["saved"]
    assert saved["status"] == "reconciled" and saved["truth"] == 45 and saved["lines"] == 3
    s = session(anon, sid, op)
    assert s["manual_count"] == 45 and s["ai_count"] == 44


def test_only_those_who_enter_tally_sheets_get_the_template(anon, roles):  # noqa: F811
    assert anon.get("/api/v1/tally/template", headers=roles["viewer"]).status_code == 403
    assert anon.get("/api/v1/tally/template").status_code == 401


def test_sheet_id_prefixes_are_the_tenants_initials():
    assert prefix_of("Bakers Inn") == "BI"
    assert prefix_of("Demo Foods (Pvt) Ltd") == "DFPL"
    assert prefix_of("") == "TS"
