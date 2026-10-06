"""Sheet uploads take Excel workbooks (.xlsx) as well as CSV: tally sheets, manifests,
the fleet register. A workbook goes through the same importer as the CSV it would
save as, so it must give the same result."""

from __future__ import annotations

import io
from datetime import date, datetime, time

from conftest import login, make_client
from openpyxl import Workbook
from test_tally_api import bay, local_window, roles, session, truck_session  # noqa: F401

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
SHEET_COLS = [
    "sheet_id", "date", "bay", "truck_plate", "direction", "start_time", "end_time",
    "route_driver", "pages", "stack_count", "total_crates", "total_on_paper", "check",
    "counted_by", "verified_by", "entered_by", "notes",
]  # fmt: skip


def _bytes(wb: Workbook) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def tally_workbook(sheet_id="BI-1", plate="ABE 2437", crates=(15, 15, 15), total=45) -> bytes:
    """Shaped like Bakers_Inn_Crate_Tally_Sheet_Template.xlsx: instructions and the paper
    form in front, real date and time cells, the green EXAMPLE row, blank rows below."""
    day, start, end = local_window()
    wb = Workbook()
    wb.active.title = "Instructions"
    wb.active.append(["Bakers Inn Crate Tally Sheet: Instructions"])
    wb.create_sheet("Tally Sheet").append(["BAKERS INN: CRATE TALLY SHEET (IVaaS POC)"])
    sheets = wb.create_sheet("Entry - Sheets")
    sheets.append(SHEET_COLS)
    as_time = lambda hhmm: time(*map(int, hhmm.split(":")))  # noqa: E731
    example = ["EXAMPLE-1", datetime(2026, 10, 12), "B1", "AEX 4821", "LOAD"]
    sheets.append([*example, time(6, 40), time(7, 5), "", 1, None, None, 124])
    real = [sheet_id, datetime.fromisoformat(day), "B1", plate, "LOAD"]
    sheets.append(
        [*real, as_time(start), as_time(end), "", 1, None, None, float(total)]
        + [None, "R. Ncube", "S. Dube"]
    )
    for _ in range(5):
        sheets.append([None] * len(SHEET_COLS))  # Excel's pre-formatted blank rows
    stacks = wb.create_sheet("Entry - Stacks")
    stacks.append(["sheet_id", "line_no", "crates", "note"])
    for i, n in enumerate(crates, start=1):
        stacks.append([sheet_id, i, float(n)])  # Excel keeps numbers as floats
    wb.create_sheet("Lists").append(["bay", "direction", "note_code"])
    return _bytes(wb)


def test_the_tally_workbook_uploaded_once_imports_like_its_two_csvs(anon, roles, bay):  # noqa: F811
    op = roles["operator"]
    sid = truck_session(anon, bay, op, crates=(15, 15, 14))  # the AI saw 44
    r = anon.post(
        f"/api/v1/tally/import?bay_id={bay['id']}",
        files={"sheets": ("Bakers_Inn_Crate_Tally_Sheet.xlsx", tally_workbook(), XLSX)},
        headers=op,
    )
    assert r.status_code == 200, r.text
    body = r.json()
    # what the CSV test asserts for the same day (test_tally_api)
    assert body["skipped"] == ["Sheets row 2: EXAMPLE-1 is the example row"]
    (saved,) = body["saved"]
    # the same as test_tally_api's CSV upload of the same day
    assert saved["status"] == "reconciled"
    assert saved["truth"] == 45 and saved["lines"] == 3 and not saved["transcription_mismatch"]
    s = session(anon, sid, op)
    assert s["manual_count"] == 45 and s["ai_count"] == 44 and s["status"] == "reconciled"


def test_a_workbook_without_the_tally_sheets_says_which_sheets_it_has(anon, roles, bay):  # noqa: F811
    wb = Workbook()
    wb.active.title = "Summary"
    wb.active.append(["truck", "crates"])
    r = anon.post(
        f"/api/v1/tally/import?bay_id={bay['id']}",
        files={"sheets": ("wrong.xlsx", _bytes(wb), XLSX)},
        headers=roles["operator"],
    )
    assert r.status_code == 422
    assert "'Entry - Sheets'" in r.text and "'Summary'" in r.text


MANIFEST_ROWS = [
    ["date", "plate", "direction", "expected", "reference", "route"],
    [date(2026, 10, 1), "ABC 1001", "LOAD", 104.0, "M-1", "Route 7"],
    [date(2026, 10, 1), "ABC 1002", "RETURN", 98.0, "M-2", "Route 9"],
]


def _manifest_csv() -> bytes:
    lines = [",".join(map(str, MANIFEST_ROWS[0]))]
    for d, p, di, e, r, rt in MANIFEST_ROWS[1:]:
        lines.append(f"{d.isoformat()},{p},{di},{int(e)},{r},{rt}")
    return ("\n".join(lines) + "\n").encode()


def _manifest_xlsx() -> bytes:
    wb = Workbook()
    wb.active.title = "Read me"  # a sheet in front without the columns: passed over
    wb.active.append(["Dispatch manifest for Bakers Inn"])
    data = wb.create_sheet("Manifest")
    for row in MANIFEST_ROWS:
        data.append(row)
    return _bytes(wb)


def test_a_manifest_workbook_imports_exactly_as_its_csv():
    results = []
    for name, payload, kind in (
        ("manifest.csv", _manifest_csv(), "text/csv"),
        ("manifest.xlsx", _manifest_xlsx(), XLSX),
    ):
        with make_client() as c:  # a fresh app for each, so neither sees the other's
            r = c.post(
                "/api/v1/manifests/import",
                files={"file": (name, payload, kind)},
                headers=login(c, "admin"),
            )
            assert r.status_code == 200, r.text
            results.append(r.json())
    assert results[0] == results[1]


def test_a_fleet_register_workbook_imports_and_reports_bad_lines():
    wb = Workbook()
    ws = wb.active
    ws.append(["Plate", "Fleet_Number", "Operator"])
    ws.append(["ABC 1234", "SL-01", "Superlink"])
    ws.append(["ABD 5678", 2.0, "Superlink"])  # a fleet number Excel stored as a number
    ws.append(["?", "SL-03", "Superlink"])
    with make_client() as c:
        admin = login(c, "admin")
        r = c.post(
            "/api/v1/fleet/import", files={"file": ("fleet.xlsx", _bytes(wb), XLSX)}, headers=admin
        )
        assert r.status_code == 200, r.text
        assert r.json() == {
            "added": 2,
            "updated": 0,
            "errors": ["line 4: a plate needs at least two letters or digits"],
        }
        fleet = {v["plate"]: v for v in c.get("/api/v1/fleet", headers=admin).json()}
        assert fleet["ABD 5678"]["fleet_number"] == "2"  # not "2.0"


def test_old_xls_and_damaged_workbooks_are_refused_with_what_to_do():
    with make_client() as c:
        admin = login(c, "admin")
        old = b"\xd0\xcf\x11\xe0" + b"\x00" * 64
        r = c.post("/api/v1/fleet/import", files={"file": ("fleet.xls", old)}, headers=admin)
        assert r.status_code == 422 and "save it as .xlsx or CSV" in r.text
        broken = b"PK\x03\x04" + b"not really a workbook"
        r = c.post("/api/v1/fleet/import", files={"file": ("fleet.xlsx", broken)}, headers=admin)
        assert r.status_code == 422 and "not a readable Excel workbook" in r.text


def test_the_blank_template_says_only_the_examples_are_filled_in(anon, roles, bay):  # noqa: F811
    wb = Workbook()
    sheets = wb.active
    sheets.title = "Entry - Sheets"
    sheets.append(SHEET_COLS)
    sheets.append(["EXAMPLE-20261012-B1-001", datetime(2026, 10, 12), "B1", "AEX 4821", "LOAD"])
    r = anon.post(
        f"/api/v1/tally/import?bay_id={bay['id']}",
        files={"sheets": ("Bakers_Inn_Crate_Tally_Sheet_Template.xlsx", _bytes(wb), XLSX)},
        headers=roles["operator"],
    )
    assert r.status_code == 422 and "Only the EXAMPLE rows are filled in" in r.text
