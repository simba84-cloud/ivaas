"""The accuracy report as a PDF and an Excel workbook: the same figures as the Accuracy
page, under the Liquid brand, with sheets that scored no load kept out of the averages."""

from __future__ import annotations

import io

from openpyxl import load_workbook
from pypdf import PdfReader
from test_sheet_uploads import XLSX
from test_tally_api import (  # noqa: F401
    bay,
    roles,
    sheets_csv,
    stacks_csv,
    truck_session,
    upload,
)


def _two_sheets(anon, roles, bay):  # noqa: F811
    """One sheet that reconciles a load (AI 44, tally 45) and one that matches none."""
    op = roles["operator"]
    truck_session(anon, bay, op, crates=(15, 15, 14))
    upload(anon, bay, op, sheets_csv(total="45"), stacks_csv(crates=(15, 15, 15)))
    upload(anon, bay, op, sheets_csv(sheet_id="BI-2", plate="ZZZ 9999", total="40"))


def _get(anon, roles, fmt):  # noqa: F811
    return anon.get("/api/v1/tally/report", params={"format": fmt}, headers=roles["viewer"])


def test_the_pdf_carries_the_figures_and_the_brand(anon, roles, bay):  # noqa: F811
    _two_sheets(anon, roles, bay)
    r = _get(anon, roles, "pdf")
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    disposition = r.headers["content-disposition"]
    assert 'filename="liquid-ivaas-bakers-inn-accuracy-' in disposition
    assert disposition.endswith('.pdf"')
    pdf = PdfReader(io.BytesIO(r.content))
    text = " ".join(p.extract_text() for p in pdf.pages)
    assert all(len(p.images) >= 1 for p in pdf.pages)  # the logo on every page
    assert "Liquid Intelligent Technologies" in text and "Accuracy report" in text
    assert "97.8%" in text and "2.2%" in text and "1 of 1" in text  # 1 - 1/45; |44-45|/45
    assert "BI-1" in text and "passed" in text
    # the unmatched sheet is listed apart, never scored as 0%
    assert "Not scored" in text and "BI-2" in text and "unmatched" in text
    assert "0.0%" not in text


def test_the_workbook_has_the_figures_as_numbers(anon, roles, bay):  # noqa: F811
    _two_sheets(anon, roles, bay)
    r = _get(anon, roles, "xlsx")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX
    assert r.headers["content-disposition"].endswith('.xlsx"')
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Every sheet"]
    assert all(len(ws._images) == 1 for ws in wb.worksheets)
    rows = list(wb["Every sheet"].iter_rows(values_only=True))
    top = next(i for i, row in enumerate(rows) if row[0] == "Sheet")
    head = rows[top]
    got = {row[0]: dict(zip(head, row, strict=False)) for row in rows[top + 1 :] if row[0]}
    scored, unmatched = got["BI-1"], got["BI-2"]
    assert scored["Tally"] == 45 and scored["AI count"] == 44 and scored["Variance"] == -1
    assert round(scored["Accuracy"], 4) == round(1 - 1 / 45, 4) and scored["Result"] == "passed"
    assert scored["Bay"] == bay["name"]
    # nothing measured is left empty, never written as 0
    assert unmatched["Accuracy"] is None and unmatched["AI count"] is None
    assert unmatched["Sheet status"] == "unmatched"


def test_with_no_sheet_scored_the_files_say_so_rather_than_zero(anon, roles, bay):  # noqa: F811
    r = _get(anon, roles, "pdf")
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert "no accuracy figure" in text and "0.0%" not in text
    wb = load_workbook(io.BytesIO(_get(anon, roles, "xlsx").content))
    values = [v for row in wb["Summary"].iter_rows(values_only=True) for v in row]
    assert any(isinstance(v, str) and "no accuracy figure" in v for v in values)


def test_json_is_still_the_default(anon, roles):  # noqa: F811
    r = anon.get("/api/v1/tally/report", headers=roles["viewer"])
    assert r.status_code == 200 and r.json()["sheets"] == 0
    assert anon.get("/api/v1/tally/report", params={"format": "pdf"}).status_code == 401


def test_one_bay_gives_that_bays_figures_like_the_page(anon, roles, bay):  # noqa: F811
    _two_sheets(anon, roles, bay)
    mine = anon.get("/api/v1/tally/report", params={"bay_id": bay["id"]}, headers=roles["viewer"])
    assert mine.json()["sheets"] == 2 and mine.json()["reconciled"] == 1
    other = "00000000-0000-0000-0000-000000000001"
    r = anon.get("/api/v1/tally/report", params={"bay_id": other}, headers=roles["viewer"])
    assert r.status_code == 404  # not someone else's bay, and not an empty report


def test_a_bays_file_says_which_bay(anon, roles, bay):  # noqa: F811
    _two_sheets(anon, roles, bay)
    r = anon.get(
        "/api/v1/tally/report",
        params={"bay_id": bay["id"], "format": "pdf"},
        headers=roles["viewer"],
    )
    slug = bay["name"].lower().replace(" ", "-")
    assert f"liquid-ivaas-bakers-inn-{slug}-accuracy-" in r.headers["content-disposition"]
    text = " ".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert f"{bay['name']} only" in text
